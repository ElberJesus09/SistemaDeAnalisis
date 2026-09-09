"""Pantallas del panel: login, inscripciones, edición, importación y reportes."""
import io
import smtplib
import tempfile
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from jinja2 import Environment, FileSystemLoader, select_autoescape

from . import auth, config, correo, copias, db, ficha, importador, pagos, sync_google
from .servidor import Respuesta, redirigir, ruta

POR_PAGINA = 50

_env = Environment(loader=FileSystemLoader(config.APP / "templates"),
                   autoescape=select_autoescape(["html"]))


def render(plantilla, pet, **ctx):
    ctx.setdefault("sesion", pet.sesion or {})
    ctx.setdefault("ciclo", config.CICLO)
    ctx.setdefault("mensaje", pet.arg("msg"))
    ctx.setdefault("google_activo", sync_google.opciones()["activo"])
    return _env.get_template(plantilla).render(**ctx)


def _nombre_completo(r):
    ape = " ".join(x for x in ((r.get("ap_paterno") or "").strip(),
                               (r.get("ap_materno") or "").strip()) if x)
    nom = (r.get("nombres") or "").strip()
    return f"{ape}, {nom}" if ape and nom else (ape or nom or "(sin nombre)")


def _ciclos():
    filas = db.q("SELECT DISTINCT ciclo FROM inscripciones ORDER BY ciclo DESC")
    lista = [f["ciclo"] for f in filas]
    if config.CICLO not in lista:
        lista.insert(0, config.CICLO)
    return lista


# =================================================================== login
@ruta("GET", "/login")
def ver_login(pet):
    aviso = auth.asegurar_admin_inicial()
    if pet.sesion:
        return redirigir("/inscripciones")
    return render("panel/login.html", pet, aviso=aviso or pet.arg("msg"),
                  error=pet.arg("error"))


@ruta("POST", "/login")
def hacer_login(pet):
    u = auth.autenticar(pet.campo("usuario"), pet.campo("clave"))
    if not u:
        return redirigir("/login?error=" + quote("Usuario o contraseña incorrectos."))
    auth.registrar(u["id"], "ingreso", u["usuario"])
    return redirigir("/inscripciones").cookie("sesion", auth.crear_sesion(u))


@ruta("GET", "/salir")
def salir(pet):
    return redirigir("/login").cookie("sesion", "", dias=0)


@ruta("GET", "/")
def inicio(pet):
    return redirigir("/inscripciones" if pet.sesion else "/login")


# =========================================================== inscripciones
def _registro_fecha(marca):
    fecha = importador.normalizar_fecha(marca)
    try:
        return datetime.strptime(fecha, "%d/%m/%Y").strftime("%Y-%m-%d")
    except ValueError:
        return ""


def _buscar(f):
    sql = ("SELECT i.id, i.ciclo, i.carrera, i.turno, i.voucher, i.fecha_matricula, i.marca_temporal,"
           " a.dni, a.nombres, a.ap_paterno, a.ap_materno, a.sexo, a.menor_edad,"
           " a.colegio, a.colegio_departamento, a.colegio_provincia,"
           " a.colegio_distrito, a.colegio_ubicacion, a.colegio_egreso,"
           " a.apo_nombres, a.apo_ap_paterno, a.apo_ap_materno,"
           " a.apo_dni, a.apo_telefono, a.apo_parentesco"
           " FROM inscripciones i JOIN alumnos a ON a.id = i.alumno_id WHERE 1=1")
    p = []
    if f["ciclo"]:
        sql += " AND i.ciclo = %s"; p.append(f["ciclo"])
    if f["carrera"]:
        sql += " AND i.carrera = %s"; p.append(f["carrera"])
    if f["q"]:
        like = f"%{f['q'].strip()}%"
        sql += (" AND (a.dni LIKE %s OR a.nombres LIKE %s"
                " OR a.ap_paterno LIKE %s OR a.ap_materno LIKE %s)")
        p += [like] * 4
    sql += " ORDER BY a.ap_paterno, a.ap_materno, a.nombres"

    filas = db.q(sql, p)
    estados = pagos.estados_por_inscripcion([r["id"] for r in filas],
                                            f["ciclo"] or config.CICLO)
    for r in filas:
        r["dia_registro"] = _registro_fecha(r["marca_temporal"])
        r["fecha_registro"] = (datetime.strptime(r["dia_registro"], "%Y-%m-%d").strftime("%d/%m/%Y")
                               if r["dia_registro"] else "Sin fecha de registro")
        r["nombre_completo"] = _nombre_completo(r)
        r["faltan"] = ficha.faltantes(r)
        r["pago"] = estados.get(r["id"], "sin_pago")
    if f.get("pago"):
        filas = [r for r in filas if r["pago"] == f["pago"]]
    if f["estado"] == "incompleto":
        filas = [r for r in filas if r["faltan"]]
    elif f["estado"] == "completo":
        filas = [r for r in filas if not r["faltan"]]
    if f.get("fecha"):
        filas = [r for r in filas if r["dia_registro"] == f["fecha"]]
    # El orden alfabetico se conserva dentro de cada dia.
    filas.sort(key=lambda r: r["dia_registro"], reverse=True)
    return filas


@ruta("GET", "/inscripciones", rol="*")
def ver_inscripciones(pet):
    f = {k: pet.arg(k) for k in ("q", "ciclo", "carrera", "estado", "pago", "fecha")}
    filas = _buscar(f)
    dias = {}
    for r in filas:
        d = dias.setdefault(r["dia_registro"], {"fecha": r["fecha_registro"], "total": 0,
                           "completo": 0, "parcial": 0, "observado": 0, "sin_pago": 0})
        d["total"] += 1
        d[r["pago"]] += 1
    try:
        pagina = max(1, int(pet.arg("p", "1")))
    except ValueError:
        pagina = 1
    paginas = max(1, -(-len(filas) // POR_PAGINA))
    pagina = min(pagina, paginas)
    trozo = filas[(pagina - 1) * POR_PAGINA: pagina * POR_PAGINA]

    base = "&".join(f"{k}={quote(v)}" for k, v in f.items() if v)
    return render("panel/inscripciones.html", pet,
                  filas=trozo, f=f, total=len(filas), dias=dias,
                  incompletos=sum(1 for r in filas if r["faltan"]),
                  obligatorios=[ficha.ETIQUETAS.get(c, c)
                                for c in ficha.OBLIGATORIOS_BASE],
                  obligatorios_apoderado=[ficha.ETIQUETAS.get(c, c)
                                          for c in ficha.OBLIGATORIOS_APODERADO],
                  masivos=MASIVOS,
                  sin_pago=sum(1 for r in filas if r["pago"] != "completo"),
                  ciclos=_ciclos(),
                  carreras=[c["nombre"] for c in
                            db.q("SELECT nombre FROM carreras ORDER BY nombre")],
                  pagina=pagina, paginas=paginas,
                  url_pagina=f"/inscripciones?{base}&p=" if base
                             else "/inscripciones?p=", seccion="inscripciones")


MASIVOS = {
    "fecha_matricula": "Fecha de matrícula",
    "turno": "Turno",
    "sede": "Sede",
    "apo_parentesco": "Parentesco del apoderado",
}


@ruta("POST", "/inscripciones/masivo", rol="*")
def llenado_masivo(pet):
    """Aplica un mismo valor a todos los registros que muestra el filtro.

    Es para los datos que el formulario no pregunta y son iguales para un
    grupo entero (la fecha de matricula, sobre todo). Solo rellena los que
    esten VACIOS: nunca pisa un dato ya puesto."""
    campo = pet.campo("campo")
    valor = pet.campo("valor").strip()
    if campo not in MASIVOS or not valor:
        return redirigir("/inscripciones", "Elige un dato y escribe un valor.")

    f = {k: pet.campo(k, "") for k in ("q", "ciclo", "carrera", "estado", "pago", "fecha")}
    filas = _buscar(f)
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    tocados = 0
    for r in filas:
        if campo == "apo_parentesco":
            fila = db.q1("SELECT a.id, a.apo_parentesco FROM alumnos a"
                         " JOIN inscripciones i ON i.alumno_id = a.id"
                         " WHERE i.id = %s", (r["id"],))
            if fila and not str(fila["apo_parentesco"] or "").strip():
                db.x("UPDATE alumnos SET apo_parentesco = %s, actualizado = %s"
                     " WHERE id = %s", (valor, ahora, fila["id"]))
                tocados += 1
        else:
            fila = db.q1(f"SELECT id, {campo} AS v FROM inscripciones WHERE id = %s",
                         (r["id"],))
            if fila and not str(fila["v"] or "").strip():
                db.x(f"UPDATE inscripciones SET {campo} = %s, actualizado = %s"
                     " WHERE id = %s", (valor, ahora, fila["id"]))
                tocados += 1

    auth.registrar(pet.sesion["id"], "masivo", campo, f"{tocados} registros = {valor}")
    base = "&".join(f"{k}={quote(v)}" for k, v in f.items() if v)
    return redirigir("/inscripciones" + ("?" + base if base else ""),
                     f"{MASIVOS[campo]}: se llenaron {tocados} registro(s) que "
                     f"estaban vacíos. Los que ya tenían dato no se tocaron.")


CAMPOS_ALUMNO = ("nombres", "ap_paterno", "ap_materno", "dni", "nacimiento", "sexo",
                 "telefono", "correo", "departamento", "provincia", "distrito",
                 "direccion", "menor_edad",
                 "colegio", "colegio_departamento", "colegio_provincia",
                 "colegio_distrito", "colegio_egreso",
                 "apo_nombres", "apo_ap_paterno", "apo_ap_materno",
                 "apo_dni", "apo_telefono", "apo_parentesco")
CAMPOS_INSC = ("carrera", "turno", "sede", "modalidad", "fecha_matricula", "voucher", "agencia",
               "fecha_pago", "medio_pago", "secuencia", "oferta")


def _fila_completa(insc_id):
    r = db.q1(
        "SELECT i.*, a.dni, a.nombres, a.ap_paterno, a.ap_materno, a.nacimiento,"
        " a.sexo, a.menor_edad, a.telefono, a.correo, a.departamento, a.provincia, a.distrito,"
        " a.direccion, a.colegio, a.colegio_ubicacion, a.colegio_egreso,"
        " a.colegio_departamento, a.colegio_provincia, a.colegio_distrito,"
        " a.apo_nombres, a.apo_ap_paterno, a.apo_ap_materno,"
        " a.apo_dni, a.apo_telefono, a.apo_parentesco, i.alumno_id"
        " FROM inscripciones i JOIN alumnos a ON a.id = i.alumno_id WHERE i.id = %s",
        (insc_id,))
    if r:
        r["nombre_completo"] = _nombre_completo(r)
    return r


@ruta("GET", r"/inscripcion/(\d+)", rol="*")
def ver_inscripcion(pet, insc_id):
    r = _fila_completa(int(insc_id))
    if not r:
        return Respuesta("<h3>Inscripción no encontrada.</h3>", 404)
    r["colegio_ubicacion_armada"] = ficha.ubicacion_colegio(r)
    faltan = [c for c in ficha.obligatorios_de(r)
              if not str((r.get("colegio_ubicacion_armada")
                          if c == "colegio_ubicacion" else r.get(c)) or "").strip()]
    return render("panel/editar.html", pet, r=r, faltan_campos=faltan,
                  faltan_nombres=[ficha.ETIQUETAS.get(c, c) for c in faltan],
                  pago=pagos.estado(r["id"], r["ciclo"]),
                  exigir_pago=config.EXIGIR_PAGO, seccion="inscripciones")


@ruta("POST", r"/inscripcion/(\d+)", rol="*")
def guardar_inscripcion(pet, insc_id):
    r = _fila_completa(int(insc_id))
    if not r:
        return Respuesta("<h3>Inscripción no encontrada.</h3>", 404)
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    va = [pet.campo(c).strip() for c in CAMPOS_ALUMNO]
    db.x(f"UPDATE alumnos SET {', '.join(c + ' = %s' for c in CAMPOS_ALUMNO)},"
         f" actualizado = %s WHERE id = %s", (*va, ahora, r["alumno_id"]))
    vi = [pet.campo(c).strip() for c in CAMPOS_INSC]
    db.x(f"UPDATE inscripciones SET {', '.join(c + ' = %s' for c in CAMPOS_INSC)},"
         f" actualizado = %s WHERE id = %s", (*vi, ahora, r["id"]))

    auth.registrar(pet.sesion["id"], "editar", f"inscripcion {insc_id}", r["dni"])
    return redirigir(f"/inscripcion/{insc_id}", "Cambios guardados.")


@ruta("GET", r"/inscripcion/(\d+)/ficha", rol="*")
def descargar_ficha(pet, insc_id):
    datos = ficha.datos_de_inscripcion(int(insc_id))
    faltan = ficha.faltantes(datos)
    if faltan:
        return redirigir(f"/inscripcion/{insc_id}",
                         "Completa primero: " + ", ".join(faltan))
    if config.EXIGIR_PAGO:
        est = pagos.estado(int(insc_id), datos["ciclo"])
        if est["situacion"] != "completo":
            return redirigir(f"/inscripcion/{insc_id}",
                             "El pago no esta verificado con el banco: "
                             + "; ".join((["Falta: " + ", ".join(est["faltan"])] if est["faltan"] else [])
                                         + est["observaciones"] or ["Falta el reporte de pagos"]))
    with tempfile.TemporaryDirectory() as tmp:
        try:
            pdf = ficha.a_pdf(datos, Path(tmp) / "ficha.pdf")
        except ValueError as exc:
            return redirigir(f"/inscripcion/{insc_id}", str(exc))
        contenido = pdf.read_bytes()
    auth.registrar(pet.sesion["id"], "ficha", f"inscripcion {insc_id}", datos["dni"])
    nombre = f"constancia_{datos['dni']}_{datos['ciclo']}.pdf"
    return Respuesta(contenido, 200, "application/pdf",
                     [("Content-Disposition", f'inline; filename="{nombre}"')])


# =============================================================== importar
@ruta("GET", "/copias", rol="admin")
def ver_copias(pet):
    return render("panel/copias.html", pet, copia=copias.estado(),
                  token=copias.token(pet.sesion), seccion="copias")


@ruta("POST", "/copias/configurar", rol="admin")
def configurar_copias(pet):
    import hmac
    if not hmac.compare_digest(pet.campo("token"), copias.token(pet.sesion)):
        return Respuesta("Solicitud no válida. Vuelve a abrir Copias de seguridad.", 403)
    try:
        copias.guardar_opciones(pet.campo('carpeta'), pet.campo('horas'), pet.campo('activo') == 'si')
    except ValueError as exc:
        return redirigir('/copias', str(exc))
    except OSError:
        return redirigir('/copias', 'No se pudo usar esa carpeta. Revisa la ruta y los permisos.')
    return redirigir('/copias', 'Configuración guardada. La nueva programación se aplica sin reiniciar.')


@ruta("POST", "/copias/crear", rol="admin")
def crear_copia(pet):
    import hmac
    if not hmac.compare_digest(pet.campo("token"), copias.token(pet.sesion)):
        return Respuesta("Solicitud no válida. Vuelve a abrir Copias de seguridad.", 403)
    resultado = copias.crear()
    return redirigir('/copias', resultado['mensaje'])


def _datos_para_correo(insc_id):
    datos = ficha.datos_de_inscripcion(int(insc_id))
    faltan = ficha.faltantes(datos)
    if faltan:
        raise ValueError("Completa primero: " + ", ".join(faltan))
    if config.EXIGIR_PAGO and pagos.estado(int(insc_id), datos["ciclo"])["situacion"] != "completo":
        raise ValueError("El pago no está verificado con el banco.")
    return datos


@ruta("GET", r"/inscripcion/(\d+)/correo", rol="*")
def ver_correo(pet, insc_id):
    try:
        datos = _datos_para_correo(insc_id)
        correo.comprobar_config()
    except LookupError:
        return Respuesta("Inscripción no encontrada.", 404)
    except ValueError as exc:
        return redirigir(f"/inscripcion/{insc_id}", str(exc))
    return render("panel/correo.html", pet, datos=datos, insc_id=insc_id,
                  remitente=config.SMTP["from_address"], asunto=correo.asunto(datos),
                  cuerpo=correo.cuerpo(datos), token=correo.crear_token(pet.sesion, insc_id),
                  seccion="inscripciones")


@ruta("POST", r"/inscripcion/(\d+)/correo", rol="*")
def enviar_correo(pet, insc_id):
    if not correo.consumir_token(pet.campo("token"), pet.sesion, insc_id):
        return Respuesta("Solicitud vencida o ya utilizada. Vuelve a abrir Enviar ficha por correo.", 409)
    try:
        datos = _datos_para_correo(insc_id)
        destinatario = correo.direccion(pet.campo("destinatario"))
        correo.comprobar_config()
        with tempfile.TemporaryDirectory() as tmp:
            pdf = ficha.a_pdf(datos, Path(tmp) / "ficha.pdf").read_bytes()
        correo.enviar_ficha(datos, destinatario, pdf)
    except LookupError:
        return Respuesta("Inscripción no encontrada.", 404)
    except ValueError as exc:
        return redirigir(f"/inscripcion/{insc_id}/correo", str(exc))
    except smtplib.SMTPAuthenticationError:
        return redirigir(f"/inscripcion/{insc_id}/correo", "Gmail rechazó el acceso. Revisa la contraseña de aplicación y los permisos de la cuenta.")
    except smtplib.SMTPRecipientsRefused:
        return redirigir(f"/inscripcion/{insc_id}/correo", "El servidor rechazó el destinatario. Revisa su correo.")
    except (OSError, smtplib.SMTPException):
        return redirigir(f"/inscripcion/{insc_id}/correo", "No se pudo confirmar el envío. Revisa Enviados en Gmail antes de intentarlo otra vez.")
    try:
        auth.registrar(pet.sesion["id"], "correo_ficha", f"inscripcion {insc_id}", destinatario)
    except Exception:
        return redirigir(f"/inscripcion/{insc_id}", "El servidor de correo aceptó la ficha, pero no se pudo guardar el historial. No repitas el envío.")
    return redirigir(f"/inscripcion/{insc_id}", f"El servidor de correo aceptó la ficha para {destinatario}. Revisa también Spam si no aparece en la bandeja de entrada.")


@ruta("GET", "/importar", rol="*")
def ver_importar(pet):
    return render("panel/importar.html", pet, resumen=None,
                  historial=_historial(), google=sync_google.estado(), seccion="importar")


@ruta("GET", "/google-sheets", rol="admin")
def ver_google_sheets(pet):
    if not sync_google.opciones()["activo"]:
        return redirigir("/importar")
    return render("panel/google_sheets.html", pet, google=sync_google.estado(),
                  token=sync_google.token(pet.sesion), seccion="importar")


@ruta("POST", "/google-sheets/sincronizar", rol="admin")
def sincronizar_google_sheets(pet):
    if not sync_google.opciones()["activo"]:
        return redirigir("/importar")
    import hmac
    if not hmac.compare_digest(pet.campo("token"), sync_google.token(pet.sesion)):
        return Respuesta("Solicitud no válida. Vuelve a abrir Google Sheets.", 403)
    mensaje = sync_google.sincronizar(pet.sesion["id"])
    return redirigir("/google-sheets", mensaje)


def _historial():
    return db.q("SELECT i.*, u.nombre AS usuario FROM importaciones i"
                " LEFT JOIN usuarios u ON u.id = i.usuario_id"
                " ORDER BY i.id DESC LIMIT 15")


@ruta("POST", "/importar", rol="*")
def hacer_importar(pet):
    subido = pet.archivos.get("archivo")
    if not subido:
        return render("panel/importar.html", pet, resumen=None,
                      historial=_historial(), seccion="importar",
                      mensaje="Elige un archivo .csv o .xlsx.",
                      clase_mensaje="mal")
    nombre, datos = subido
    ciclo = pet.campo("ciclo").strip() or config.CICLO
    with tempfile.TemporaryDirectory() as tmp:
        ruta_tmp = Path(tmp) / Path(nombre).name
        ruta_tmp.write_bytes(datos)
        try:
            # El tipo se deduce de los encabezados: el reporte del banco trae
            # DOCUMENTO y COD_PAGO; el del formulario, DNI y NOMBRE COMPLETO.
            enc, _ = importador.leer_archivo_crudo(ruta_tmp)
            es_pagos = pagos.es_archivo_de_pagos(enc)
            modulo = pagos if es_pagos else importador
            res = modulo.importar(ruta_tmp, pet.sesion["id"], ciclo, nombre)
            res["tipo"] = "pagos" if es_pagos else "inscripciones"
        except Exception as e:
            return render("panel/importar.html", pet, resumen=None,
                          historial=_historial(), seccion="importar",
                          mensaje=f"No se pudo importar: {e}", clase_mensaje="mal")
    res["archivo"] = nombre
    auth.registrar(pet.sesion["id"], "importar", nombre,
                   f"{res['nuevas']} nuevas, {res['actualizadas']} actualizadas")
    return render("panel/importar.html", pet, resumen=res,
                  historial=_historial(), seccion="importar")


# =============================================================== reportes
def _agrupar(campo, ciclo, tabla="i"):
    filas = db.q(
        f"SELECT {tabla}.{campo} AS clave, COUNT(*) AS n"
        " FROM inscripciones i JOIN alumnos a ON a.id = i.alumno_id"
        " WHERE (%s = '' OR i.ciclo = %s)"
        f" GROUP BY {tabla}.{campo} ORDER BY n DESC", (ciclo, ciclo))
    total = sum(f["n"] for f in filas) or 1
    for f in filas:
        f["pct"] = round(f["n"] * 100 / total)
    return filas


def _recaudacion(ciclo):
    """Cuanto entro, por concepto y por agencia, segun el reporte del banco."""
    filas = db.q("SELECT cod_pago, concepto, COUNT(*) AS n, SUM(importe) AS soles"
                 " FROM pagos WHERE ciclo = %s GROUP BY cod_pago, concepto"
                 " ORDER BY cod_pago", (ciclo,))
    por_agencia = db.q("SELECT agencia AS clave, COUNT(*) AS n,"
                       " SUM(importe) AS soles FROM pagos WHERE ciclo = %s"
                       " GROUP BY agencia ORDER BY soles DESC", (ciclo,))
    for f in filas + por_agencia:
        f["soles"] = pagos._num(f["soles"])
    return {"conceptos": filas, "agencias": por_agencia,
            "total": sum(f["soles"] for f in filas),
            "pagos": sum(f["n"] for f in filas)}


@ruta("GET", "/reportes", rol="*")
def ver_reportes(pet):
    ciclos = _ciclos()
    ciclo = pet.arg("ciclo") or (ciclos[0] if ciclos else config.CICLO)
    filas = _buscar({"q": "", "ciclo": ciclo, "carrera": "", "estado": "", "pago": ""})
    incompletos = sum(1 for r in filas if r["faltan"])
    pagados = sum(1 for r in filas if r["pago"] == "completo")
    kpi = {"total": len(filas), "incompletos": incompletos,
           "completos": len(filas) - incompletos,
           "carreras": len({r["carrera"] for r in filas if r["carrera"]}),
           "pagados": pagados, "sin_pagar": len(filas) - pagados}
    return render("panel/reportes.html", pet, kpi=kpi, ciclos=ciclos,
                  ciclo_sel=ciclo, seccion="reportes",
                  caja=_recaudacion(ciclo),
                  por_carrera=_agrupar("carrera", ciclo),
                  por_turno=_agrupar("turno", ciclo),
                  por_agencia=_agrupar("agencia", ciclo),
                  por_departamento=_agrupar("departamento", ciclo, "a"),
                  por_fecha=_agrupar("fecha_pago", ciclo))


@ruta("GET", "/reportes/excel", rol="*")
def reportes_excel(pet):
    import openpyxl

    ciclo = pet.arg("ciclo") or config.CICLO
    wb = openpyxl.Workbook()

    hoja = wb.active
    hoja.title = "Inscripciones"
    cab = ["DNI", "Apellidos y nombres", "Carrera", "Ciclo", "Modalidad", "Turno", "Voucher",
           "Agencia", "Fecha pago", "Estado", "Faltan", "Pago"]
    hoja.append(cab)
    for r in _buscar({"q": "", "ciclo": ciclo, "carrera": "", "estado": "", "pago": ""}):
        det = db.q1("SELECT agencia, fecha_pago, modalidad FROM inscripciones WHERE id = %s",
                    (r["id"],)) or {}
        hoja.append([r["dni"], r["nombre_completo"], r["carrera"], r["ciclo"],
                     det.get("modalidad", ""),
                     r["turno"], r["voucher"], det.get("agencia", ""),
                     det.get("fecha_pago", ""),
                     "Incompleto" if r["faltan"] else "Completo",
                     ", ".join(r["faltan"]),
                     {"completo": "Pagado", "parcial": "Pago parcial", "observado": "Pago observado"}
                     .get(r["pago"], "Sin pago")])

    hp = wb.create_sheet("Pagos")
    hp.append(["DNI", "Nombre en el banco", "Voucher", "Ult. 7", "Concepto",
               "Importe", "Fecha", "Hora", "Agencia", "Caja", "Aplicado"])
    for p in db.q("SELECT p.*, a.dni AS dni_alumno FROM pagos p"
                  " LEFT JOIN inscripciones i ON i.id = p.inscripcion_id"
                  " LEFT JOIN alumnos a ON a.id = i.alumno_id"
                  " WHERE p.ciclo = %s ORDER BY p.fecha_pago, p.hora", (ciclo,)):
        hp.append([p["dni"], p["nombre_banco"], p["voucher"], p["voucher_clave"],
                   p["concepto"] or p["cod_pago"], float(pagos._num(p["importe"])),
                   p["fecha_pago"], p["hora"], p["agencia"], p["caja"],
                   "Si" if p["inscripcion_id"] else "Sin inscripcion"])

    caja = _recaudacion(ciclo)
    hr = wb.create_sheet("Recaudacion")
    hr.append(["Concepto", "Pagos", "Soles"])
    for f in caja["conceptos"]:
        hr.append([f["concepto"] or f["cod_pago"], f["n"], float(f["soles"])])
    hr.append([])
    hr.append(["Agencia", "Pagos", "Soles"])
    for f in caja["agencias"]:
        hr.append([f["clave"] or "(sin dato)", f["n"], float(f["soles"])])
    hr.append([])
    hr.append(["TOTAL", caja["pagos"], float(caja["total"])])

    for titulo, datos in (("Por carrera", _agrupar("carrera", ciclo)),
                          ("Por turno", _agrupar("turno", ciclo)),
                          ("Por agencia", _agrupar("agencia", ciclo)),
                          ("Por departamento", _agrupar("departamento", ciclo, "a")),
                          ("Por fecha de pago", _agrupar("fecha_pago", ciclo))):
        h = wb.create_sheet(titulo[:31])
        h.append([titulo.replace("Por ", "").capitalize(), "Inscritos", "%"])
        for f in datos:
            h.append([f["clave"] or "(sin dato)", f["n"], f["pct"]])

    for h in wb.worksheets:
        for col in h.columns:
            ancho = max((len(str(c.value or "")) for c in col), default=10)
            h.column_dimensions[col[0].column_letter].width = min(46, ancho + 3)

    buf = io.BytesIO()
    wb.save(buf)
    return Respuesta(
        buf.getvalue(), 200,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        [("Content-Disposition",
          f'attachment; filename="reportes_{ciclo}.xlsx"')])


# =============================================================== usuarios
@ruta("GET", "/usuarios", rol="admin")
def ver_usuarios(pet):
    return render("panel/usuarios.html", pet, seccion="usuarios",
                  usuarios=db.q("SELECT id, usuario, nombre, rol, activo"
                                " FROM usuarios ORDER BY usuario"))


@ruta("POST", "/usuarios", rol="admin")
def crear_usuario(pet):
    try:
        auth.crear_usuario(pet.campo("usuario"), pet.campo("nombre"),
                           pet.campo("clave"), pet.campo("rol", "secretaria"))
    except Exception as e:
        return redirigir("/usuarios", f"No se pudo crear: {e}")
    return redirigir("/usuarios", "Usuario creado.")


@ruta("POST", "/usuarios/clave", rol="admin")
def cambiar_clave(pet):
    auth.cambiar_clave(int(pet.campo("id")), pet.campo("clave"))
    return redirigir("/usuarios", "Contraseña actualizada.")


@ruta("POST", "/usuarios/activo", rol="admin")
def activar_usuario(pet):
    db.x("UPDATE usuarios SET activo = %s WHERE id = %s",
         (int(pet.campo("activo")), int(pet.campo("id"))))
    return redirigir("/usuarios", "Estado actualizado.")


# ================================================================== pagos
@ruta("GET", "/pagos", rol="*")
def ver_pagos(pet):
    ciclos = _ciclos()
    ciclo = pet.arg("ciclo") or (ciclos[0] if ciclos else config.CICLO)
    pagos.asegurar_tarifario(ciclo)
    q = pet.arg("q").strip()
    solo = pet.arg("solo")

    sql = ("SELECT p.*, a.dni AS dni_alumno, a.ap_paterno, a.ap_materno, a.nombres"
           " FROM pagos p"
           " LEFT JOIN inscripciones i ON i.id = p.inscripcion_id"
           " LEFT JOIN alumnos a ON a.id = i.alumno_id"
           " WHERE p.ciclo = %s")
    par = [ciclo]
    if solo == "huerfanos":
        sql += " AND p.inscripcion_id IS NULL"
    elif solo == "aplicados":
        sql += " AND p.inscripcion_id IS NOT NULL"
    if q:
        sql += " AND (p.dni LIKE %s OR p.voucher LIKE %s OR p.nombre_banco LIKE %s)"
        par += [f"%{q}%"] * 3
    sql += " ORDER BY p.fecha_pago DESC, p.hora DESC"
    filas = db.q(sql, par)
    for r in filas:
        r["alumno"] = _nombre_completo(r) if r.get("dni_alumno") else ""

    total = sum(pagos._num(r["importe"]) for r in filas)
    sueltos = db.q1("SELECT COUNT(*) AS n FROM pagos WHERE inscripcion_id IS NULL"
                    " AND ciclo = %s", (ciclo,))
    return render("panel/pagos.html", pet, filas=filas, ciclo_sel=ciclo,
                  ciclos=ciclos, q=q, solo=solo, total=total,
                  huerfanos=sueltos["n"] if sueltos else 0,
                  tarifario=pagos.tarifario(ciclo), seccion="pagos")


@ruta("POST", "/pagos/conciliar", rol="*")
def reconciliar(pet):
    ciclo = pet.campo("ciclo") or config.CICLO
    r = pagos.conciliar(ciclo)
    auth.registrar(pet.sesion["id"], "conciliar", ciclo, f"{r['aplicados']} pagos")
    return redirigir(f"/pagos?ciclo={quote(ciclo)}",
                     f"{r['aplicados']} pago(s) aplicados. "
                     f"{r['huerfanos']} sin inscripción.")


# =============================================================== tarifario
@ruta("GET", "/tarifario", rol="admin")
def ver_tarifario(pet):
    ciclos = _ciclos()
    ciclo = pet.arg("ciclo") or config.CICLO
    pagos.asegurar_tarifario(ciclo)
    return render("panel/tarifario.html", pet, ciclos=ciclos, ciclo_sel=ciclo,
                  editar=db.q1("SELECT * FROM conceptos_pago WHERE ciclo = %s AND id = %s",
                               (ciclo, pet.arg("editar"))) if pet.arg("editar").isdigit() else None,
                  conceptos=db.q("SELECT * FROM conceptos_pago WHERE ciclo = %s"
                                 " ORDER BY codigo", (ciclo,)), seccion="tarifario")


@ruta("POST", "/tarifario", rol="admin")
def guardar_tarifario(pet):
    ciclo = pet.campo("ciclo") or config.CICLO
    try:
        codigo = pagos.codigo_pago(pet.campo("codigo"))
        importe = pagos.importe_valido(pet.campo("importe"))
        nombre = pet.campo("nombre").strip()
        if not nombre or len(nombre) > 80:
            raise ValueError("Indica un nombre de hasta 80 caracteres.")
    except ValueError as exc:
        return redirigir(f"/tarifario?ciclo={quote(ciclo)}", str(exc))
    datos = (nombre, str(importe), 1,
             1 if pet.campo("obligatorio") else 0)
    ya = db.q1("SELECT id FROM conceptos_pago WHERE ciclo = %s AND codigo = %s",
               (ciclo, codigo))
    if ya:
        db.x("UPDATE conceptos_pago SET nombre = %s, importe = %s, fijo = %s,"
             " obligatorio = %s, activo = 1 WHERE id = %s", (*datos, ya["id"]))
    else:
        db.x("INSERT INTO conceptos_pago (ciclo, codigo, nombre, importe, fijo,"
             " obligatorio, activo) VALUES (%s, %s, %s, %s, %s, %s, 1)",
             (ciclo, codigo, *datos))
    auth.registrar(pet.sesion["id"], "tarifario", f"{ciclo}/{codigo}", datos[0])
    return redirigir(f"/tarifario?ciclo={quote(ciclo)}", "Tarifario guardado.")


@ruta("POST", "/tarifario/quitar", rol="admin")
def quitar_concepto(pet):
    ciclo = pet.campo("ciclo") or config.CICLO
    db.x("UPDATE conceptos_pago SET activo = 0 WHERE id = %s",
         (int(pet.campo("id")),))
    return redirigir(f"/tarifario?ciclo={quote(ciclo)}", "Concepto desactivado.")
