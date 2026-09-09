"""
Prueba de extremo a extremo del sistema completo.

Levanta el servidor real sobre una base SQLite temporal y recorre el mismo
camino que hara la secretaria: entrar, importar el archivo del formulario,
ver el listado, completar los datos que faltan, emitir la ficha en PDF y
descargar los reportes.

    python tests/test_sistema.py
"""
import csv
import io
import os
import sys
import tempfile
import threading
import urllib.error
import urllib.request
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

TMP = Path(tempfile.mkdtemp(prefix="cpu_test_"))
os.environ["CPU_TEST"] = "1"

from app import auth, config, db, ficha, importador, pagos, servidor, web  # noqa: E402

fallos = []
hechas = 0


def ok(condicion, titulo, detalle=""):
    global hechas
    hechas += 1
    if condicion:
        print(f"  OK   {titulo}")
    else:
        print(f"  FALLA {titulo} {detalle}")
        fallos.append(titulo)


# ------------------------------------------------------------ datos de prueba
ENCABEZADOS = [
    "Marca temporal", "Dirección de correo electrónico", "MEDIO DE PAGO",
    "OFERTAS", "NUMERO DE VOUCHER ", "FECHA DE PAGO", "AGENCIA DE PAGO",
    "NUMERO DE SECUENCIA", "FECHA DE PAGO",
    "¿A qué carrera profesional deseas postular? ", "DNI", "NOMBRE COMPLETO",
    "APELLIDO PATERNO", "APELLIDO MATERNO", "FECHA DE NACIMIENTO", "CELULAR ",
    "CORREO ELECTRONICO", "DEPARTAMENTO", "  PROVINCIA ", "DISTRITO", "DIRECCIÓN",
]


def fila(dni, nombres, pat, mat, carrera, voucher, dep="LAMBAYEQUE"):
    return ["12/06/2026 9:15:03", "form@gmail.com", "Banco de la Nación", "Regular",
            voucher, "12/06/2026", "0230", "998877", "",
            carrera, dni, nombres, pat, mat, "01/07/2007", "937635827",
            "alumno@gmail.com", dep, "CHICLAYO", "JOSE LEONARDO ORTIZ",
            "CALLE LOS ROBLES 245"]


FILAS = [
    fila("70123456", "ANA LUCIA", "PEREZ", "GOMEZ", "DERECHO", "052135"),
    fila("70123457", "CARLOS ANDRES", "QUISPE", "TORRES", "ADMINISTRACIÓN", "052136"),
    fila("70123458", "MARIA JOSE", "SANCHEZ", "QUIROZ", "DERECHO", "052137", "PIURA"),
    fila("70123457", "CARLOS ANDRES", "QUISPE", "TORRES", "ADMINISTRACIÓN", "052136"),
    fila("", "SIN DOCUMENTO", "X", "Y", "DERECHO", "052138"),
]


def escribir_csv(ruta):
    with Path(ruta).open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(ENCABEZADOS)
        w.writerows(FILAS)


def escribir_xlsx(ruta):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(ENCABEZADOS)
    for f in FILAS[:3]:
        ws.append(f)
    wb.save(ruta)


ENC_NUEVO = [
    "Marca temporal", "Dirección de correo electrónico", "MEDIO DE PAGO", "SEDE",
    "TURNO", "NUMERO DE VOUCHER ", "FECHA DE PAGO", "AGENCIA DE PAGO",
    "NUMERO DE SECUENCIA", "FECHA DE PAGO",
    "¿A qué carrera profesional deseas postular? ", "DNI", "NOMBRE COMPLETO",
    "APELLIDO PATERNO", "APELLIDO MATERNO", "GENERO", "FECHA DE NACIMIENTO",
    "CELULAR ", "CORREO ELECTRONICO", "DEPARTAMENTO", "  PROVINCIA ", "DISTRITO",
    "DIRECCIÓN", "ERES MENOR DE EDAD", "NOMBRE DEL APODERADO",
    "APELLIDO PATERNO  DEL APODERADO", "APELLIDO MATERNO   DEL APODERADO",
    "DNI DEL APODERADO", "TELEFONO DEL APODERADO", "NOMBRE DEL COLEGIO",
    "DEPARTAMENTO ", "PROVINCIA", "DISTRITO", "AÑO DE EGRESO",
]


def valores_nuevo(dni="71000001", menor="SI", genero="Femenino"):
    """Una fila del formulario nuevo. Ojo: la terna del alumno es distinta de
    la del colegio, justamente para detectar si se pisan."""
    apo = ["CARLOS", "RAMIREZ", "SOTO", "16700001", "979000111"] if menor == "SI" \
        else ["", "", "", "", ""]
    return ["20/06/2026 8:30:00", "form@gmail.com", "Banco de la Nación",
            "Sede Central", "Tarde", "8200111-1", "", "0230", "998001",
            "20/06/2026", "MEDICINA HUMANA", dni, "ANA MARIA", "TORRES", "DIAZ",
            genero, "05/05/2008", "955000111", "ana@gmail.com",
            "LAMBAYEQUE", "CHICLAYO", "JOSE LEONARDO ORTIZ", "CALLE UNION 100",
            menor, *apo, "I.E. NUESTRA SENORA",
            "PIURA", "SULLANA", "BELLAVISTA", "2025"]


def escribir_nuevo(ruta):
    with Path(ruta).open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(ENC_NUEVO)
        w.writerow(valores_nuevo())
        w.writerow(valores_nuevo("71000002", "NO", "Masculino"))


ENC_BANCO = ["NRO", "COD.", "COD_ALUMNO", "DOCUMENTO", "VOUCHER", "CODIGO_TDOC",
             "NOMBRE_TDOC", "SITUACION", "COD_PAGO", "CONCEPTO_PAGO",
             "APELLIDOS_NOMBRES", "CUENTA", "FECHA_PAGO", "HORA", "IMPORTE_S/.",
             "CAJ.", "AGE."]

# dni, voucher, cod_pago, importe, nombre
PAGOS_BANCO = [
    ("70123456", "8100987", "00001097", 750, "PEREZ GOMEZ ANA LUCIA"),
    ("70123456", "8100123", "00001096", 200, "#PEREZ GOMEZ ANA LUCIA"),
    ("70123457", "8101234", "00001096", 150, "QUISPE TORRES CARLOS"),   # pagó de menos
    ("88888888", "8109999", "00001096", 200, "AJENO SIN FORMULARIO"),   # sin inscripción
    ("88888888", "8109998", "00001097", 750, "AJENO SIN FORMULARIO"),
    ("70123456", "8100987", "00001097", 750, "PEREZ GOMEZ ANA LUCIA"),  # repetido
    ("", "8100000", "00001096", 200, "SIN DOCUMENTO"),                  # se omite
]


def escribir_banco(ruta):
    """El reporte del banco viene separado por tabuladores."""
    with Path(ruta).open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(ENC_BANCO)
        for i, (dni, vou, cod, imp, nombre) in enumerate(PAGOS_BANCO, start=1):
            doc = (dni + "0000000") if dni else ""
            w.writerow([i, "004", "-", doc, vou, "01", "DNI", "00090009", cod,
                        "*** DESCCONOCIDO ***", nombre, "0301029403",
                        "18-06-2026", "10:0%d:00" % (i % 10), imp, "1033", "0231"])


# ------------------------------------------------------------------ cliente
class Cliente:
    """Cliente HTTP minimo que recuerda la cookie de sesion."""

    def __init__(self, base):
        self.base = base
        self.cookie = ""

    def pedir(self, ruta, datos=None, cabeceras=None, seguir=True):
        req = urllib.request.Request(self.base + ruta, data=datos,
                                     headers=cabeceras or {})
        if self.cookie:
            req.add_header("Cookie", self.cookie)

        class SinRedireccion(urllib.request.HTTPRedirectHandler):
            def redirect_request(self, *a, **k):
                return None

        abridor = urllib.request.build_opener() if seguir else \
            urllib.request.build_opener(SinRedireccion)
        try:
            r = abridor.open(req)
            estado, cuerpo, cabs = r.status, r.read(), r.headers
        except urllib.error.HTTPError as e:
            estado, cuerpo, cabs = e.code, e.read(), e.headers
        for c in cabs.get_all("Set-Cookie") or []:
            if c.startswith("sesion="):
                self.cookie = c.split(";")[0]
        return estado, cuerpo, cabs

    def get(self, ruta, seguir=True):
        return self.pedir(ruta, seguir=seguir)

    def post(self, ruta, campos, seguir=True):
        from urllib.parse import urlencode
        return self.pedir(ruta, urlencode(campos).encode(),
                          {"Content-Type": "application/x-www-form-urlencoded"},
                          seguir)

    def subir(self, ruta, campos, archivo):
        """POST multipart, tal como lo manda un navegador."""
        frontera = "----WebKitFormBoundary7MA4YWxkTrZu0gW"
        buf = io.BytesIO()
        for k, v in campos.items():
            buf.write(f"--{frontera}\r\nContent-Disposition: form-data; "
                      f'name="{k}"\r\n\r\n{v}\r\n'.encode())
        nombre, contenido = archivo
        buf.write(f"--{frontera}\r\nContent-Disposition: form-data; "
                  f'name="archivo"; filename="{nombre}"\r\n'
                  "Content-Type: application/octet-stream\r\n\r\n".encode())
        buf.write(contenido)
        buf.write(f"\r\n--{frontera}--\r\n".encode())
        return self.pedir(ruta, buf.getvalue(),
                          {"Content-Type":
                           f"multipart/form-data; boundary={frontera}"})


# ------------------------------------------------------------------ pruebas
def main():
    print("\n1. Base de datos y esquema")
    db.usar_sqlite(TMP / "prueba.db")
    db.crear_esquema()
    tablas = {t["name"] for t in db.q(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    ok({"alumnos", "inscripciones", "usuarios", "importaciones",
        "auditoria", "carreras"} <= tablas, "se crean todas las tablas", tablas)

    ok(db.verificar_tablas() == "", "la estructura recien creada se valida")

    print("\n1b. Aviso si la base ya tiene una tabla de otro proyecto")
    db.x("ALTER TABLE usuarios RENAME TO usuarios_ok")
    db.x("CREATE TABLE usuarios (id INTEGER PRIMARY KEY, login TEXT)")
    problema = db.verificar_tablas()
    ok("usuarios" in problema and "clave" in problema,
       "detecta la tabla ajena y dice que columnas faltan", problema[:90])
    db.x("DROP TABLE usuarios")
    db.x("ALTER TABLE usuarios_ok RENAME TO usuarios")
    ok(db.verificar_tablas() == "", "tras restaurarla vuelve a validar")

    print("\n2. Contraseñas y sesiones")
    h = auth.hash_clave("secreta123")
    ok(auth.verificar_clave("secreta123", h), "la contraseña correcta valida")
    ok(not auth.verificar_clave("otra", h), "una contraseña falsa no valida")
    ok("secreta123" not in h, "la contraseña no se guarda en claro")
    aviso = auth.asegurar_admin_inicial()
    ok("admin" in aviso, "se crea el usuario admin inicial")
    galleta = auth.crear_sesion({"id": 1, "usuario": "admin",
                                 "nombre": "Administrador", "rol": "admin"})
    ok(auth.leer_sesion(galleta)["rol"] == "admin", "la sesión firmada se lee")
    ok(auth.leer_sesion(galleta[:-3] + "xyz") is None,
       "una sesión manipulada se rechaza")

    print("\n3. Importador (CSV)")
    csv_ruta = TMP / "respuestas.csv"
    escribir_csv(csv_ruta)
    r = importador.importar(csv_ruta, None, "2026-II")
    ok(r["nuevas"] == 3, "importa las 3 filas válidas", r)
    ok(r["omitidas"] == 2, "omite el DNI repetido y el vacío", r)
    ok(len(db.q("SELECT * FROM alumnos")) == 3, "3 alumnos en la base")

    a = db.q1("SELECT * FROM alumnos WHERE dni = %s", ("70123456",))
    ok(a["nombres"] == "ANA LUCIA" and a["ap_paterno"] == "PEREZ",
       "nombres y apellidos separados")
    ok(a["nacimiento"] == "01/07/2007", "la fecha queda en dd/mm/aaaa")
    i = db.q1("SELECT * FROM inscripciones WHERE alumno_id = %s", (a["id"],))
    ok(i["voucher"] == "052135", "el voucher se guarda")
    ok(i["fecha_pago"] == "12/06/2026",
       "la columna FECHA DE PAGO duplicada toma el valor no vacío", i["fecha_pago"])
    ok(i["ciclo"] == "2026-II", "el ciclo se aplica")

    print("\n4. Reimportar el mismo archivo no duplica")
    r2 = importador.importar(csv_ruta, None, "2026-II")
    ok(r2["nuevas"] == 0 and r2["actualizadas"] == 3, "no crea duplicados", r2)
    ok(len(db.q("SELECT * FROM inscripciones")) == 3, "siguen siendo 3 inscripciones")

    print("\n5. Importador (XLSX)")
    xlsx_ruta = TMP / "respuestas.xlsx"
    escribir_xlsx(xlsx_ruta)
    r3 = importador.importar(xlsx_ruta, None, "2026-II")
    ok(r3["filas"] == 3 and r3["nuevas"] == 0, "lee el Excel y reconoce lo existente", r3)

    print("\n6. Campos faltantes")
    insc = db.q1("SELECT id FROM inscripciones ORDER BY id LIMIT 1")
    datos = ficha.datos_de_inscripcion(insc["id"])
    faltan = ficha.faltantes(datos)
    ok("Sexo" in faltan and "Turno" in faltan,
       "detecta los datos que el formulario no recoge", faltan)
    ok(datos["direccion"].endswith("LAMBAYEQUE"),
       "la dirección se arma con distrito, provincia y departamento", datos["direccion"])

    print("\n7. Servidor y login")
    srv = servidor.crear(0)
    puerto = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    c = Cliente(f"http://127.0.0.1:{puerto}")

    est, cuerpo, _ = c.get("/login")
    ok(est == 200 and b"Sistema de Inscripciones" in cuerpo, "carga la pantalla de entrada")
    est, _, _ = c.get("/inscripciones", seguir=False)
    ok(est == 303, "sin sesión redirige al login", est)
    c.post("/login", {"usuario": "admin", "clave": "mala"}, seguir=False)
    ok(not c.cookie, "una contraseña incorrecta no entrega sesión")
    c.post("/login", {"usuario": "admin", "clave": "admin"}, seguir=False)
    ok(bool(c.cookie), "la contraseña correcta entrega sesión")

    print("\n8. Panel")
    est, cuerpo, _ = c.get("/inscripciones")
    ok(est == 200 and b"PEREZ" in cuerpo, "el listado muestra a los alumnos")
    ok("Sin pago" in cuerpo.decode(), "el listado muestra la situación de pago")
    ok(b"Faltan" in cuerpo, "marca los registros incompletos")
    est, cuerpo, _ = c.get("/inscripciones?estado=completo")
    ok(b"No hay inscripciones" in cuerpo, "el filtro de completos no devuelve nada aún")
    est, cuerpo, _ = c.get("/inscripciones?q=QUISPE")
    ok(b"QUISPE" in cuerpo and b"PEREZ" not in cuerpo, "la búsqueda filtra")

    est, cuerpo, _ = c.get(f"/inscripcion/{insc['id']}")
    ok(est == 200 and b"Datos del apoderado" in cuerpo, "abre la ficha de edición")

    print("\n9. La ficha se bloquea si faltan datos")
    from urllib.parse import unquote as _unq
    est, _, cabs = c.get(f"/inscripcion/{insc['id']}/ficha", seguir=False)
    destino = _unq(cabs.get("Location", ""))
    ok(est == 303 and "Completa primero" in destino,
       "no deja emitir la ficha incompleta", destino)

    print("\n10. Completar datos y emitir la ficha")
    completo = {
        "nombres": "ANA LUCIA", "ap_paterno": "PEREZ", "ap_materno": "GOMEZ",
        "dni": "70123456", "nacimiento": "01/07/2007", "sexo": "Femenino",
        "telefono": "937635827", "correo": "ana@gmail.com",
        "departamento": "LAMBAYEQUE", "provincia": "CHICLAYO",
        "distrito": "JOSE LEONARDO ORTIZ", "direccion": "CALLE LOS ROBLES 245",
        "menor_edad": "Sí",
        "colegio": "I.E. SAN JOSE", "colegio_departamento": "LAMBAYEQUE",
        "colegio_provincia": "CHICLAYO", "colegio_distrito": "CHICLAYO",
        "colegio_egreso": "2024", "apo_nombres": "JOSE", "apo_ap_paterno": "PEREZ",
        "apo_ap_materno": "RAMOS",
        "apo_dni": "16712345", "apo_telefono": "979111222", "apo_parentesco": "Padre",
        "carrera": "DERECHO", "turno": "Mañana", "sede": config.SEDE,
        "fecha_matricula": "12/06/2026", "voucher": "052135", "agencia": "0230",
        "fecha_pago": "12/06/2026", "medio_pago": "Banco de la Nación",
        "secuencia": "998877", "oferta": "Regular",
    }
    c.post(f"/inscripcion/{insc['id']}", completo, seguir=False)
    guardado = db.q1("SELECT sexo, colegio FROM alumnos WHERE dni = %s", ("70123456",))
    ok(guardado["sexo"] == "Femenino" and guardado["colegio"] == "I.E. SAN JOSE",
       "los datos completados se guardan")
    ok(not ficha.faltantes(ficha.datos_de_inscripcion(insc["id"])),
       "ya no quedan campos obligatorios vacíos")

    est, pdf, cabs = c.get(f"/inscripcion/{insc['id']}/ficha", seguir=False)
    ok(est == 200 and pdf[:4] == b"%PDF", "descarga la ficha en PDF", est)
    ok("application/pdf" in cabs.get("Content-Type", ""), "se envía como PDF")
    salida = TMP / "ficha.pdf"
    salida.write_bytes(pdf)
    import pdfplumber
    with pdfplumber.open(salida) as doc:
        paginas = len(doc.pages)
        texto = doc.pages[0].extract_text()
    ok(paginas == 1, f"la ficha ocupa UNA sola hoja (salieron {paginas})")
    ok("PEREZ" in texto and "2026-II" in texto and "I.E. SAN JOSE" in texto,
       "la ficha trae los datos del alumno y el ciclo correcto")
    ok("LAMBAYEQUE / CHICLAYO / CHICLAYO" in texto,
       "la ubicación del colegio se arma con las tres partes")
    ok("JOSE PEREZ RAMOS" in texto,
       "el apoderado se arma con nombre y los dos apellidos")

    print("\n9b. Formulario nuevo: 34 columnas y encabezados repetidos")
    ok(importador.normalizar_sexo("MASCULINO") == "Masculino"
       and importador.normalizar_sexo("Femenino") == "Femenino"
       and importador.normalizar_sexo("F") == "Femenino"
       and importador.normalizar_sexo("Mujer") == "Femenino",
       "GENERO se normaliza a Masculino/Femenino")
    ok(importador.normalizar_si_no("SI") == "Sí"
       and importador.normalizar_si_no("no") == "No",
       "ERES MENOR DE EDAD se normaliza a Sí/No")

    fila = importador._mapear(ENC_NUEVO, valores_nuevo())
    ok(fila["departamento"] == "LAMBAYEQUE" and fila["provincia"] == "CHICLAYO"
       and fila["distrito"] == "JOSE LEONARDO ORTIZ",
       "la primera terna de DEPARTAMENTO/PROVINCIA/DISTRITO es del alumno", fila)
    ok(fila["colegio_departamento"] == "PIURA"
       and fila["colegio_provincia"] == "SULLANA"
       and fila["colegio_distrito"] == "BELLAVISTA",
       "la segunda terna es del COLEGIO, no se pisa con la del alumno", fila)
    ok(fila["sede"] == "Sede Central" and fila["turno"] == "Tarde",
       "SEDE y TURNO llegan del formulario")
    ok(fila["apo_ap_paterno"] == "RAMIREZ" and fila["apo_ap_materno"] == "SOTO",
       "los apellidos del apoderado no se confunden con los del alumno", fila)
    ok(fila["fecha_pago"] == "20/06/2026",
       "la FECHA DE PAGO repetida sigue tomando el valor no vacío")

    nuevo_csv = TMP / "formulario_nuevo.csv"
    escribir_nuevo(nuevo_csv)
    rn = importador.importar(nuevo_csv, None, "2026-II")
    ok(rn["nuevas"] == 2, "importa el archivo del formulario nuevo", rn)

    men = db.q1("SELECT * FROM alumnos WHERE dni = %s", ("71000001",))
    ok(men["sexo"] == "Femenino" and men["menor_edad"] == "Sí",
       "guarda género y si es menor de edad", men["sexo"])
    ok(men["colegio_distrito"] == "BELLAVISTA" and men["distrito"] == "JOSE LEONARDO ORTIZ",
       "el distrito del colegio y el del alumno quedan separados en la base")

    im = db.q1("SELECT i.* FROM inscripciones i WHERE i.alumno_id = %s", (men["id"],))
    fm = ficha.datos_de_inscripcion(im["id"])
    ok(fm["colegio_ubicacion"] == "PIURA / SULLANA / BELLAVISTA",
       "la ficha arma la ubicación del colegio", fm["colegio_ubicacion"])
    ok(fm["apo_nombres"] == "CARLOS RAMIREZ SOTO",
       "la ficha arma el nombre del apoderado", fm["apo_nombres"])
    ok(ficha.faltantes(men) == ["Fecha de matrícula"] or
       "Fecha de matrícula" in ficha.faltantes(men),
       "al menor solo le falta la fecha de matrícula", ficha.faltantes(men))

    may = db.q1("SELECT * FROM alumnos WHERE dni = %s", ("71000002",))
    ok(may["menor_edad"] == "No", "el segundo alumno queda como mayor de edad")
    ok("Apoderado" not in ficha.faltantes(may),
       "a un mayor de edad NO se le exige apoderado", ficha.faltantes(may))
    men_sin_apo = dict(men, apo_nombres="", apo_ap_paterno="", apo_ap_materno="",
                       apo_dni="", apo_telefono="")
    ok("Apoderado" in ficha.faltantes(men_sin_apo),
       "pero a un menor sin apoderado sí se le exige",
       ficha.faltantes(men_sin_apo))
    ok("Apoderado" not in ficha.faltantes(men),
       "y al menor que sí lo trae del formulario no se le pide nada")

    print("\n9c. Migración de columnas nuevas sobre una base vieja")
    import tempfile as _tmp
    from app import db as _db
    ruta_vieja = Path(_tmp.mkdtemp()) / "vieja.db"
    guardada = (_db._MOTOR, _db._RUTA_SQLITE)
    _db.usar_sqlite(ruta_vieja)
    _db.x("CREATE TABLE alumnos (id INTEGER PRIMARY KEY AUTOINCREMENT,"
          " dni VARCHAR(12) NOT NULL UNIQUE, nombres VARCHAR(120) NOT NULL DEFAULT '',"
          " creado DATETIME NOT NULL, actualizado DATETIME NOT NULL)")
    _db.crear_esquema()
    cols = _db._columnas_actuales("alumnos")
    ok({"menor_edad", "colegio_departamento", "colegio_provincia",
        "colegio_distrito", "apo_ap_paterno", "apo_ap_materno"} <= cols,
       "una base creada con la versión anterior recibe las columnas nuevas")
    ok(_db.verificar_tablas() == "", "y queda validada")
    _db.usar_sqlite(TMP / "prueba.db")      # volver a la base de la prueba

    print("\n10b. Regla de Págalo.pe para el voucher")
    ok(pagos.clave_voucher("1234567-1") == "1234567",
       "del formulario se toman los 7 dígitos antes del guion")
    ok(pagos.clave_voucher("81234567") == "1234567",
       "del banco se toman los últimos 7 dígitos")
    ok(pagos.clave_voucher("1234567-1") == pagos.clave_voucher("81234567"),
       "1234567-1 y 81234567 coinciden — el caso de tu ejemplo")
    ok(pagos.clave_voucher("3392126") == "3392126", "un voucher de 7 dígitos se respeta")
    ok(pagos.clave_voucher("") == "" and pagos.clave_voucher(None) == "",
       "un voucher vacío no genera clave")
    ok(pagos.clave_voucher("0052135-2") == "0052135", "conserva los ceros de la izquierda")
    ok(pagos.normalizar_documento("619426920000000") == "61942692",
       "el DNI se recorta del documento relleno con ceros")
    ok(pagos.normalizar_documento("72451803") == "72451803", "un DNI normal no se toca")
    ok(str(pagos.a_decimal("1,250.50")) == "1250.50", "los importes con coma se leen bien")

    print("\n10c. Importar el reporte de pagos del banco")
    banco = TMP / "banco.csv"
    escribir_banco(banco)
    enc, _ = importador.leer_archivo_crudo(banco)
    ok(pagos.es_archivo_de_pagos(enc), "reconoce que es un archivo de pagos")
    ok(not pagos.es_archivo_de_pagos(ENCABEZADOS),
       "y que el del formulario NO lo es")

    rp = pagos.importar(banco, None, "2026-II")
    ok(rp["nuevas"] == 5, "entran los 5 pagos distintos", rp)
    ok(rp["omitidas"] == 1, "omite la fila sin DNI", rp)
    ok(rp["repetidas"] == 1, "detecta la línea repetida dentro del archivo", rp)
    ok(rp["importe"] == 2050,
       "el total no cuenta dos veces la línea repetida", str(rp["importe"]))
    ok(len(db.q("SELECT * FROM pagos")) == 5, "solo se guardaron 5 pagos")

    rp2 = pagos.importar(banco, None, "2026-II")
    ok(rp2["nuevas"] == 0 and rp2["actualizadas"] == 5,
       "reimportar el reporte no duplica pagos", rp2)

    p1 = db.q1("SELECT * FROM pagos WHERE dni = %s AND cod_pago = %s",
               ("70123456", "00001096"))
    ok(p1 is not None and str(p1["importe"]) in ("200", "200.0", "200.00"),
       "guarda el pago de S/200 del primer alumno")
    ok(p1["fecha_pago"] == "18/06/2026", "la fecha del banco pasa a dd/mm/aaaa")
    ok(p1["concepto"] == "Derecho de inscripción",
       "toma el nombre del tarifario porque el banco lo manda en blanco",
       p1["concepto"])
    ok(p1["nombre_banco"].startswith("PEREZ"), "limpia el # del nombre del banco")
    ok(p1["inscripcion_id"] == insc["id"], "el pago queda amarrado a su inscripción")

    print("\n10d. Estado de pago y no reutilización")
    est = pagos.estado(insc["id"], "2026-II", "8100987-1")
    ok(est["situacion"] == "completo", "con los dos pagos queda completo", est["situacion"])
    ok(str(est["total"]) in ("950", "950.0", "950.00"), "suma S/950", str(est["total"]))
    ok(est["voucher_ok"] is True, "el voucher declarado coincide con un pago suyo")

    est_mal = pagos.estado(insc["id"], "2026-II", "9999999-1")
    ok(est_mal["voucher_ok"] is False, "avisa si el voucher declarado no es suyo")

    otra = db.q1("SELECT i.id FROM inscripciones i JOIN alumnos a ON a.id = i.alumno_id"
                 " WHERE a.dni = %s", ("70123457",))
    est2 = pagos.estado(otra["id"], "2026-II")
    ok(est2["situacion"] == "parcial" and est2["faltan"] == ["Pensión del ciclo"],
       "con un solo pago queda parcial y dice cuál falta", est2)

    tercero = db.q1("SELECT i.id FROM inscripciones i JOIN alumnos a ON a.id = i.alumno_id"
                    " WHERE a.dni = %s", ("70123458",))
    ok(pagos.estado(tercero["id"], "2026-II")["situacion"] == "sin_pago",
       "sin pagos queda en sin_pago")

    antes = db.q1("SELECT inscripcion_id FROM pagos WHERE id = %s", (p1["id"],))
    pagos.conciliar("2026-II")
    despues = db.q1("SELECT inscripcion_id FROM pagos WHERE id = %s", (p1["id"],))
    ok(antes["inscripcion_id"] == despues["inscripcion_id"],
       "un pago ya aplicado no se mueve a otro alumno al reconciliar")

    sueltos = pagos.huerfanos("2026-II")
    ok(len(sueltos) == 2 and {p["dni"] for p in sueltos} == {"88888888"},
       "los pagos de quien no llenó el formulario quedan listados aparte",
       len(sueltos))

    print("\n10e. El importe fijo se controla")
    det = {d["codigo"]: d for d in pagos.estado(otra["id"], "2026-II")["detalle"]}
    ok("pagó S/ 150" in det["00001096"]["aviso"],
       "avisa si pagó menos de los S/200 del derecho fijo", det["00001096"]["aviso"])

    print("\n11. Importar desde el navegador (multipart)")
    est, cuerpo, _ = c.subir("/importar", {"ciclo": "2026-II"},
                             ("respuestas.csv", csv_ruta.read_bytes()))
    ok(est == 200 and b"fila(s) le" in cuerpo, "la subida de archivo funciona", est)
    ok(b"omitida" in cuerpo, "informa de las filas omitidas")
    ok("respuestas del formulario" in cuerpo.decode(),
       "detecta solo que es el archivo del formulario")

    est, cuerpo, _ = c.subir("/importar", {"ciclo": "2026-II"},
                             ("banco.csv", banco.read_bytes()))
    ok("reporte de pagos del banco" in cuerpo.decode(),
       "y que el otro es el reporte del banco", est)

    est, cuerpo, _ = c.get("/pagos?ciclo=2026-II")
    ok(est == 200 and b"3392126" not in cuerpo and b"8100987" in cuerpo,
       "la pantalla de pagos lista los pagos importados")
    est, cuerpo, _ = c.get("/pagos?ciclo=2026-II&solo=huerfanos")
    ok(b"88888888" in cuerpo, "puede filtrar los pagos sin inscripción")

    est, cuerpo, _ = c.get("/tarifario")
    ok(est == 200 and b"00001096" in cuerpo, "el admin ve el tarifario")
    c.post("/tarifario", {"ciclo": "2026-II", "codigo": "00001098",
                          "nombre": "Carné", "importe": "20", "fijo": "1"},
           seguir=False)
    ok(db.q1("SELECT id FROM conceptos_pago WHERE codigo = %s", ("00001098",))
       is not None, "puede agregar un concepto al tarifario")
    db.x("DELETE FROM conceptos_pago WHERE codigo = %s", ("00001098",))
    ok(len(db.q("SELECT * FROM importaciones")) >= 1, "queda registro de la importación")

    print("\n11b. Llenado masivo para registros sin fecha")
    # Simula registros antiguos sin marca temporal ni fecha de matricula.
    db.x("UPDATE inscripciones SET fecha_matricula='' WHERE id <> %s", (insc["id"],))
    antes = db.q1("SELECT fecha_matricula FROM inscripciones WHERE id = %s",
                  (insc["id"],))["fecha_matricula"]
    ok(bool(antes), "el primer alumno ya tenía fecha de matrícula puesta a mano")
    c.post("/inscripciones/masivo",
           {"campo": "fecha_matricula", "valor": "01/07/2026", "ciclo": "2026-II"},
           seguir=False)
    despues = db.q1("SELECT fecha_matricula FROM inscripciones WHERE id = %s",
                    (insc["id"],))["fecha_matricula"]
    ok(despues == antes, "el llenado masivo NO pisa un dato que ya existía")
    otros = db.q("SELECT fecha_matricula FROM inscripciones WHERE id <> %s",
                 (insc["id"],))
    ok(all(o["fecha_matricula"] == "01/07/2026" for o in otros),
       "y sí rellena todos los que estaban vacíos", otros)
    ok(db.q1("SELECT id FROM auditoria WHERE accion = %s", ("masivo",)) is not None,
       "el llenado masivo queda en la auditoría")

    print("\n12. Reportes")
    est, cuerpo, _ = c.get("/reportes?ciclo=2026-II")
    ok(est == 200 and b"Inscritos por carrera" in cuerpo, "la pantalla de reportes carga")
    ok(b"DERECHO" in cuerpo, "agrupa por carrera")
    est, xls, cabs = c.get("/reportes/excel?ciclo=2026-II")
    ok(est == 200 and xls[:2] == b"PK", "descarga el Excel de reportes")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xls))
    ok("Inscripciones" in wb.sheetnames and "Por carrera" in wb.sheetnames,
       "el Excel trae las hojas esperadas", wb.sheetnames)

    print("\n13. Permisos por rol")
    est, _, _ = c.get("/usuarios")
    ok(est == 200, "el admin entra a usuarios")
    auth.crear_usuario("secre", "Secretaría", "clave123", "secretaria")
    c2 = Cliente(f"http://127.0.0.1:{puerto}")
    c2.post("/login", {"usuario": "secre", "clave": "clave123"}, seguir=False)
    est, _, _ = c2.get("/inscripciones")
    ok(est == 200, "secretaría entra al listado")
    est, cuerpo, _ = c2.get("/usuarios")
    ok(est == 403, "secretaría NO entra a usuarios", est)

    srv.shutdown()

    print("\n" + "=" * 58)
    if fallos:
        print(f"{len(fallos)} de {hechas} comprobaciones FALLARON:")
        for f in fallos:
            print("   -", f)
        return 1
    print(f"Las {hechas} comprobaciones pasaron.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
